# 透明枠テンプレートのWordに、表・裏のPNGを10面ずつ配置する。
import sys, json
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

tools = Path(sys.argv[1])      # テンプレート/自動処理
assets = Path(sys.argv[2])     # 書き出し先（表.png・裏.png・書き出し情報.json）
out = Path(sys.argv[3])        # 出力する .docx

info = json.loads((assets / '書き出し情報.json').read_text(encoding='utf-8'))
school = (info.get('school') or '教室名未設定').strip()

d = Document(tools / 'A4-10面_透明枠.docx')
assert len(d.tables) == 2 and all(len(t.rows) == 5 and len(t.columns) == 2 for t in d.tables)
for n, t in enumerate(d.tables):
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]; p.paragraph_format.line_spacing = Pt(1)
            run = p.add_run()
            inline = run.add_picture(str(assets / ('表.png' if n == 0 else '裏.png')), width=Mm(91), height=Mm(55))._inline
            # セルの段落に固定して、行の高さによるはみ出しをなくす。
            anchor = OxmlElement('wp:anchor')
            for k, v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'0','locked':'1','layoutInCell':'1','allowOverlap':'1'}.items():
                anchor.set(k, v)
            e = OxmlElement('wp:simplePos'); e.set('x','0'); e.set('y','0'); anchor.append(e)
            for axis, rel in [('H','column'), ('V','paragraph')]:
                e = OxmlElement('wp:position'+axis); e.set('relativeFrom', rel)
                offset = OxmlElement('wp:posOffset'); offset.text = '0'; e.append(offset); anchor.append(e)
            anchor.append(inline.find(qn('wp:extent')))
            anchor.append(OxmlElement('wp:wrapNone'))
            for tag in ['wp:docPr','wp:cNvGraphicFramePr','a:graphic']:
                e = inline.find(qn(tag))
                if e is not None: anchor.append(e)
            anchor.find(qn('wp:docPr')).set('descr', '友達紹介カード ' + ('表' if n == 0 else '裏'))
            inline.getparent().replace(inline, anchor)

d.core_properties.title = '友達紹介カード ' + school + ' A4 10面'
d.core_properties.comments = '表裏2ページ。A4、倍率100%、両面は短辺とじ。枠線なし。元HTMLの保存済み写真と配置を反映。'
out.parent.mkdir(parents=True, exist_ok=True)
d.save(out)
print('Wordを作成しました: ' + str(out))
